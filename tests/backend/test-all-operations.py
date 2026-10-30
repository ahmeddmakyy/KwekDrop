"""Test all operations in KwekDrop Catalog, Convert wheel pairs, and preview paths.

Prints a formatted table:
operation | input | result | seconds | note
and summary:
N passed, M failed
Exits with code 1 on any failure, 0 on all pass.
"""
from pathlib import Path
import sys, subprocess, json, hashlib, math, wave, time, tempfile, shutil
from PIL import Image, ImageDraw, ImageFont
import pillow_heif, numpy as np
from pypdf import PdfReader
from docx import Document
import pymupdf as fitz

ROOT = Path(__file__).resolve().parents[2]
PYTHON = ROOT / 'outputs/KwekDrop/runtime/python/python.exe'
WORKER = ROOT / 'src/KwekDrop/backend/worker.py'
FFMPEG = ROOT / 'outputs/KwekDrop/runtime/ffmpeg/ffmpeg.exe'
FFPROBE = ROOT / 'outputs/KwekDrop/runtime/ffmpeg/ffprobe.exe'
SEVEN = ROOT / 'outputs/KwekDrop/runtime/7zip/7z.exe'

sys.path.insert(0, str(ROOT / 'src/KwekDrop/backend'))
import archives

# Create clean temporary directory for test fixtures and outputs
TEMP_DIR = Path(tempfile.mkdtemp(prefix='kwekdrop-all-ops-'))
FIXTURES = TEMP_DIR / 'fixtures'
OUTPUTS = TEMP_DIR / 'outputs'
FIXTURES.mkdir(parents=True, exist_ok=True)
OUTPUTS.mkdir(parents=True, exist_ok=True)

def run_proc(args, timeout=120):
    res = subprocess.run(list(map(str, args)), capture_output=True, timeout=timeout, creationflags=0x08000000)
    if res.returncode != 0:
        err = res.stderr.decode('utf-8', 'replace')[-1200:]
        raise RuntimeError(err or f"Process failed with exit code {res.returncode}")
    return res.stdout

def probe(path):
    out = run_proc([FFPROBE, '-v', 'error', '-show_format', '-show_streams', '-of', 'json', str(path)])
    return json.loads(out)

def execute_job(name, paths, action, params=None):
    req_file = OUTPUTS / f"{name}.req.json"
    res_file = OUTPUTS / f"{name}.res.json"
    req_file.write_text(json.dumps({
        'Paths': [str(p) for p in paths],
        'Action': action,
        'Parameters': {**(params or {}), '_outputDir': str(OUTPUTS)}
    }, ensure_ascii=False), encoding='utf-8')

    run_proc([PYTHON, '-I', WORKER, '--job', str(req_file), str(res_file)])
    data = json.loads(res_file.read_text(encoding='utf-8'))
    errors = [f.get('Error') for f in data.get('Files', []) if f.get('Error')]
    if errors:
        raise RuntimeError('; '.join(errors))
    return [Path(f['Output']) for f in data.get('Files', [])]

def execute_preview(name, path, action, params=None, mode='--preview'):
    req_file = OUTPUTS / f"{name}.prev_req.json"
    res_file = OUTPUTS / f"{name}.prev_res.json"
    prev_folder = OUTPUTS / f"{name}_prev_dir"
    prev_folder.mkdir(parents=True, exist_ok=True)

    req_file.write_text(json.dumps({
        'Paths': [str(path)],
        'Action': action,
        'Parameters': params or {}
    }, ensure_ascii=False), encoding='utf-8')

    run_proc([PYTHON, '-I', WORKER, mode, str(req_file), str(res_file), str(prev_folder)])
    data = json.loads(res_file.read_text(encoding='utf-8'))
    if 'Error' in data:
        raise RuntimeError(data['Error'])
    prev_file = data.get('Preview')
    if not prev_file or not Path(prev_file).exists():
        raise RuntimeError(f"Preview file not found: {prev_file}")
    return Path(prev_file)

# ----------------- CREATE FIXTURES -----------------
print("Creating test fixtures in", FIXTURES)
fixtures = {}

# Images
img = Image.new('RGB', (160, 120), '#FFFFFF')
dr = ImageDraw.Draw(img)
dr.rectangle((0, 0, 79, 59), fill='#FF5200')
dr.rectangle((80, 0, 159, 59), fill='#22C55E')
dr.rectangle((0, 60, 79, 119), fill='#3B82F6')
dr.text((85, 75), 'KWEK', fill='#12151A')

img.save(FIXTURES / 'image.png', format='PNG')
img.save(FIXTURES / 'image.jpg', format='JPEG')
img.save(FIXTURES / 'image.webp', format='WEBP')
img.save(FIXTURES / 'image.bmp', format='BMP')
pillow_heif.register_heif_opener()
img.save(FIXTURES / 'image.heic', format='HEIF')

fixtures['png'] = FIXTURES / 'image.png'
fixtures['jpg'] = FIXTURES / 'image.jpg'
fixtures['webp'] = FIXTURES / 'image.webp'
fixtures['bmp'] = FIXTURES / 'image.bmp'
fixtures['heic'] = FIXTURES / 'image.heic'

# Image with OCR-friendly text
ocr_img = Image.new('RGB', (400, 200), '#FFFFFF')
ocr_dr = ImageDraw.Draw(ocr_img)
ocr_dr.text((40, 80), 'KWEKDROP OCR TEST', fill='#000000')
ocr_img.save(FIXTURES / 'ocr_sample.png', format='PNG')
fixtures['ocr_img'] = FIXTURES / 'ocr_sample.png'

# SVG & GIF
svg_text = '<svg xmlns="http://www.w3.org/2000/svg" width="160" height="120"><rect width="160" height="120" fill="white"/><rect width="80" height="60" fill="#FF5200"/><circle cx="120" cy="80" r="20" fill="#3B82F6"/></svg>'
(FIXTURES / 'image.svg').write_text(svg_text, encoding='utf-8')
fixtures['svg'] = FIXTURES / 'image.svg'

gif_frames = [Image.new('RGB', (160, 120), c) for c in ('#FF5200', '#22C55E', '#3B82F6')]
gif_frames[0].save(FIXTURES / 'movie.gif', save_all=True, append_images=gif_frames[1:], duration=300, loop=0)
fixtures['gif'] = FIXTURES / 'movie.gif'

# Audio (3s stereo 44.1kHz sine wave)
rate = 44100
dur = 3.0
t_samples = np.arange(int(rate * dur)) / rate
left = np.sin(t_samples * 2 * math.pi * 440) * 0.4
right = np.sin(t_samples * 2 * math.pi * 660) * 0.4
pcm_data = (np.stack([left, right], axis=1) * 32767).astype('<i2')

wav_path = FIXTURES / 'sound.wav'
with wave.open(str(wav_path), 'wb') as wf:
    wf.setnchannels(2)
    wf.setsampwidth(2)
    wf.setframerate(rate)
    wf.writeframes(pcm_data.tobytes())
fixtures['wav'] = wav_path

mp3_path = FIXTURES / 'sound.mp3'
run_proc([FFMPEG, '-y', '-i', str(wav_path), '-c:a', 'libmp3lame', '-b:a', '160k', str(mp3_path)])
fixtures['mp3'] = mp3_path

# Video (3s 160x120 video with test audio)
mp4_path = FIXTURES / 'movie.mp4'
run_proc([FFMPEG, '-y', '-f', 'lavfi', '-i', 'testsrc=size=160x120:rate=24:duration=3.0', '-i', str(wav_path), '-shortest', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', '-c:a', 'aac', str(mp4_path)])
fixtures['mp4'] = mp4_path

silent_path = FIXTURES / 'silent.mp4'
run_proc([FFMPEG, '-y', '-f', 'lavfi', '-i', 'testsrc=size=160x120:rate=24:duration=3.0', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', str(silent_path)])
fixtures['silent'] = silent_path

# PDF (3 pages with text and image)
pdf_path = FIXTURES / 'document.pdf'
pdoc = fitz.open()
for p in range(3):
    page = pdoc.new_page(width=300, height=400)
    page.insert_text((40, 50), f"Page {p+1} KwekDrop Test Document")
    page.insert_image(fitz.Rect(40, 80, 200, 200), filename=str(fixtures['png']))
pdoc.set_metadata({'title': 'KwekDrop Sample', 'author': 'Antigravity'})
pdoc.save(str(pdf_path))
pdoc.close()
fixtures['pdf'] = pdf_path

# DOCX
docx_path = FIXTURES / 'document.docx'
doc = Document()
doc.add_heading('KwekDrop Document Test', 0)
doc.add_paragraph('First paragraph testing Word processing.')
doc.save(str(docx_path))
fixtures['docx'] = docx_path

# Text & Subtitles
txt_path = FIXTURES / 'notes.txt'
txt_path.write_text('KwekDrop test note.\nArabic text: تجربة تطبيق كويك دروب\nEnglish text: Desktop file tools.', encoding='utf-8')
fixtures['txt'] = txt_path

srt_path = FIXTURES / 'caption.srt'
srt_path.write_text('1\n00:00:00,000 --> 00:00:01,000\nHello World\n\n2\n00:00:01,000 --> 00:00:02,500\nKwekDrop Subtitle\n', encoding='utf-8')
fixtures['srt'] = srt_path

# Archives
payload = FIXTURES / 'archive_payload'
(payload / 'sub').mkdir(parents=True, exist_ok=True)
(payload / 'sub' / 'data.txt').write_text('Nested archive content.', encoding='utf-8')
(payload / 'readme.txt').write_text('Root readme file.', encoding='utf-8')

zip_path = FIXTURES / 'bundle.zip'
archives.pack(payload, zip_path, 'zip')
fixtures['zip'] = zip_path

# Store source hashes
source_hashes = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in fixtures.values() if p.is_file()}

# ----------------- TEST OPERATIONS -----------------
records = []

def run_test(op_name, inputs_list, action, params=None, preview_mode=None):
    t0 = time.time()
    input_str = ', '.join(p.name for p in inputs_list) if isinstance(inputs_list, list) else Path(inputs_list).name
    paths = inputs_list if isinstance(inputs_list, list) else [inputs_list]
    
    try:
        # Run execution
        outs = execute_job(op_name, paths, action, params)
        out = outs[0]
        assert out.exists()
        size = out.stat().st_size if out.is_file() else sum(f.stat().st_size for f in out.iterdir())
        assert size > 0

        # Validate by format
        suf = out.suffix.lower()
        if suf in ('.png', '.jpg', '.jpeg', '.webp', '.bmp', '.heic'):
            with Image.open(out) as im:
                im.load()
                assert im.width > 0 and im.height > 0
        elif suf in ('.mp4', '.mov', '.mkv', '.webm', '.avi'):
            pr = probe(out)
            assert any(s['codec_type'] == 'video' for s in pr['streams'])
        elif suf in ('.mp3', '.wav', '.flac', '.m4a', '.ogg', '.opus', '.aiff', '.wma', '.m4r'):
            pr = probe(out)
            assert any(s['codec_type'] == 'audio' for s in pr['streams'])
        elif suf == '.pdf':
            doc = fitz.open(str(out))
            assert len(doc) > 0
            doc.close()
        elif suf == '.docx':
            d = Document(str(out))
            assert len(d.paragraphs) > 0
        elif suf == '.zip':
            import zipfile
            assert zipfile.is_zipfile(str(out))
        elif suf == '.txt':
            assert len(out.read_text(encoding='utf-8', errors='ignore')) > 0

        prev_note = ""
        # If preview requested
        if preview_mode:
            p_out = execute_preview(op_name, paths[0], action, params, mode=preview_mode)
            assert p_out.exists() and p_out.stat().st_size > 0
            prev_note = f" + preview ok ({p_out.name})"

        dt = time.time() - t0
        records.append({
            'operation': op_name,
            'input': input_str,
            'result': 'PASS',
            'seconds': dt,
            'note': f"out: {out.name} ({size}B){prev_note}",
            'passed': True,
            'output_path': out
        })
        return out
    except Exception as ex:
        dt = time.time() - t0
        records.append({
            'operation': op_name,
            'input': input_str,
            'result': 'FAIL',
            'seconds': dt,
            'note': str(ex)[:180],
            'passed': False,
            'output_path': None
        })
        return None

print("\nRunning Catalog Operations...")

# IMAGE TOOLS
run_test("cropImage", fixtures['png'], "cropImage", {'x':'10','y':'10','width':'64','height':'48'}, preview_mode='--preview')
run_test("rotateImage", fixtures['png'], "rotateImage", {'angle':'90','flip':'horizontal'}, preview_mode='--preview')
run_test("resizeImage", fixtures['png'], "resizeImage", {'mode':'edge','edge':'64'})
run_test("editImage", fixtures['png'], "editImage", {'exposure':'0.5','saturation':'1.2','clarity':'10'}, preview_mode='--preview')
run_test("frameImage", fixtures['png'], "frameImage", {'canvasWidth':'400','canvasHeight':'300','padding':'30'}, preview_mode='--preview')
run_test("watermark", fixtures['png'], "watermark", {'text':'Kwek','opacity':'80'}, preview_mode='--preview')
run_test("redactImage", fixtures['png'], "redactImage", {'regions':json.dumps([{'x':10,'y':10,'width':30,'height':20,'style':'solid'}])}, preview_mode='--preview')
run_test("compressImage", fixtures['jpg'], "compress", {'quality':'70','maxEdge':'120'})
run_test("removeMetadataImage", fixtures['jpg'], "removeMetadata", {'remove':'true'})
loc_jpg = run_test("locationImageSet", fixtures['jpg'], "location", {'mode':'set','coordinates':'30.0444, 31.2357'})
if loc_jpg:
    run_test("locationImageStrip", loc_jpg, "location", {'mode':'strip'})
run_test("ocrImage", fixtures['ocr_img'], "ocrImage", {'output':'txt'})
run_test("makeIcon", fixtures['png'], "makeIcon", {'iconKind':'ico'})
run_test("createPDF", [fixtures['png'], fixtures['jpg']], "createPDF", {})
run_test("createCollage", [fixtures['png'], fixtures['jpg']], "createCollage", {'canvasWidth':'400','canvasHeight':'300','layout':'grid'})
run_test("createAnimation", [fixtures['png'], fixtures['jpg']], "createAnimation", {'format':'gif','seconds':'0.5'})
run_test("packArchiveImage", fixtures['jpg'], "packArchive", {'format':'zip'})

# VIDEO TOOLS
run_test("trimVideo", fixtures['mp4'], "trimVideo", {'start':'0.5','end':'2.5'}, preview_mode='--playback')
run_test("splitVideo", fixtures['mp4'], "splitVideo", {'splitPoints':'1.0,2.0'})
run_test("joinVideos", [fixtures['mp4'], fixtures['silent']], "joinVideos", {})
run_test("changeVideoSpeed", fixtures['mp4'], "changeVideoSpeed", {'speed':'1.5'}, preview_mode='--playback')
run_test("videoSnapshots", fixtures['mp4'], "videoSnapshots", {'times':'0.5,1.5,2.5'})
run_test("videoSnapshotsSheet", fixtures['mp4'], "videoSnapshots", {'sheet':'true','sheetCount':'4','sheetColumns':'2'})
run_test("videoEffects", fixtures['mp4'], "videoEffects", {'fadeIn':'0.5','fadeOut':'0.5'})
run_test("cropVideo", fixtures['mp4'], "cropVideo", {'x':'10','y':'10','width':'64','height':'48'}, preview_mode='--playback')
run_test("rotateVideo", fixtures['mp4'], "rotateVideo", {'angle':'180'}, preview_mode='--playback')
run_test("videoSettings", fixtures['mp4'], "videoSettings", {'height':'72','fps':'15'})
run_test("watermarkVideo", fixtures['mp4'], "watermark", {'text':'KwekVideo'})
run_test("redactVideo", fixtures['mp4'], "redactVideo", {'regions':json.dumps([{'x':10,'y':10,'width':30,'height':30,'start':0,'end':2,'style':'solid'}])}, preview_mode='--playback')
run_test("muteVideo", fixtures['mp4'], "muteVideo", {}, preview_mode='--playback')
run_test("subtitlesAudio", fixtures['mp4'], "subtitlesAudio", {'subtitleFile':str(fixtures['srt'])})
run_test("compressVideo", fixtures['mp4'], "compress", {'quality':'32','maxEdge':'120'})
run_test("removeMetadataVideo", fixtures['mp4'], "removeMetadata", {'remove':'true'})
loc_mp4 = run_test("locationVideoSet", fixtures['mp4'], "location", {'mode':'set','coordinates':'30.0444, 31.2357'})
if loc_mp4:
    run_test("locationVideoStrip", loc_mp4, "location", {'mode':'strip'})
run_test("videoToGif", fixtures['mp4'], "videoToGif", {'width':'80','fps':'8'})
run_test("packArchiveVideo", fixtures['mp4'], "packArchive", {'format':'zip'})

# AUDIO TOOLS
run_test("trimAudio", fixtures['wav'], "trimAudio", {'start':'0.5','end':'2.0','removeSilence':'false'}, preview_mode='--playback')
run_test("redactAudio", fixtures['wav'], "redactAudio", {'ranges':json.dumps([{'start':0.5,'end':1.5}])}, preview_mode='--playback')
run_test("joinAudio", [fixtures['wav'], fixtures['wav']], "joinAudio", {})
run_test("normalizeAudio", fixtures['wav'], "normalizeAudio", {'loudness':'-16'}, preview_mode='--playback')
run_test("audioChannels", fixtures['wav'], "audioChannels", {'channels':'mono'}, preview_mode='--playback')
run_test("audioEffects", fixtures['wav'], "audioEffects", {'gainDb':'3','fadeIn':'0.5','fadeOut':'0.5','noiseReduction':'true'}, preview_mode='--playback')
run_test("compressAudio", fixtures['mp3'], "compress", {'bitrate':'96'})
run_test("removeMetadataAudio", fixtures['mp3'], "removeMetadata", {'remove':'true'})
run_test("audioToVideo", fixtures['wav'], "audioToVideo", {'aspect':'landscape','waveStyle':'mirror','color':'orange'}, preview_mode='--playback')
run_test("ringtone", fixtures['wav'], "ringtone", {'start':'0','length':'2'})
run_test("packArchiveAudio", fixtures['mp3'], "packArchive", {'format':'zip'})

# PDF / DOCUMENT TOOLS
run_test("splitPDF", fixtures['pdf'], "splitPDF", {'pagesPerFile':'1'})
run_test("mergePDF", [fixtures['pdf'], fixtures['pdf']], "mergePDF", {})
run_test("organizePDF", fixtures['pdf'], "organizePDF", {'pageOrder':'3,2,1','rotation':'90'}, preview_mode='--preview')
run_test("extractPdfImages", fixtures['pdf'], "extractPdfImages", {})
run_test("pdfNumbers", fixtures['pdf'], "pdfNumbers", {'numberFormat':'{n} / {total}'}, preview_mode='--preview')
run_test("pdfPassword", fixtures['pdf'], "pdfPassword", {'mode':'add','newPassword':'secret'})
run_test("watermarkPDF", fixtures['pdf'], "watermark", {'text':'CONFIDENTIAL'})
run_test("redactPDF", fixtures['pdf'], "redactPDF", {'regions':json.dumps([{'page':1,'x':40,'y':40,'width':120,'height':20}])}, preview_mode='--preview')
run_test("compressPDF", fixtures['pdf'], "compress", {'quality':'65'})
run_test("removeMetadataPDF", fixtures['pdf'], "removeMetadata", {'remove':'true'})
run_test("ocrPDF", fixtures['pdf'], "ocrPDF", {'output':'txt'})

# OFFICE TOOLS
run_test("officePdf", fixtures['docx'], "officePdf", {})
run_test("officeMergePDF", [fixtures['docx']], "officeMergePDF", {})
run_test("removeMetadataDocx", fixtures['docx'], "removeMetadata", {'remove':'true'})

# ARCHIVE TOOLS
run_test("extractArchive", fixtures['zip'], "extractArchive", {})
run_test("packArchive", fixtures['txt'], "packArchive", {'format':'zip'})

# CONVERT WHEEL PAIRS
convert_pairs = [
    # Image targets
    ('png', 'jpg'), ('jpg', 'webp'), ('webp', 'png'), ('png', 'pdf'), ('png', 'bmp'),
    # Video targets
    ('mp4', 'mov'), ('mp4', 'mkv'), ('mp4', 'webm'), ('mp4', 'gif'), ('mp4', 'mp3'),
    # Audio targets
    ('wav', 'mp3'), ('mp3', 'flac'), ('wav', 'm4a'), ('wav', 'ogg'),
    # Document targets
    ('pdf', 'docx'), ('pdf', 'txt'), ('txt', 'pdf'), ('txt', 'md'), ('srt', 'vtt'),
    # Archive targets
    ('zip', 'tar'), ('zip', 'gz')
]
for src_ext, dst_ext in convert_pairs:
    if src_ext in fixtures:
        run_test(f"convert_{src_ext}_to_{dst_ext}", fixtures[src_ext], f"convert:{dst_ext}", {})

# VERIFY SOURCE FIXTURES UNTOUCHED
sources_ok = True
for path_str, orig_hash in source_hashes.items():
    cur_hash = hashlib.sha256(Path(path_str).read_bytes()).hexdigest()
    if cur_hash != orig_hash:
        sources_ok = False
        print(f"ERROR: Source file was modified: {path_str}")

# PRINT TABLE
print("\n" + "="*110)
print(f"{'Operation':32s} | {'Input':22s} | {'Result':7s} | {'Sec':6s} | {'Note'}")
print("="*110)
for r in records:
    print(f"{r['operation']:32s} | {r['input'][:22]:22s} | {r['result']:7s} | {r['seconds']:5.2f}s | {r['note'][:40]}")
print("="*110)

passed_cnt = sum(1 for r in records if r['passed'])
failed_cnt = sum(1 for r in records if not r['passed'])

print(f"\nSummary: {passed_cnt} passed, {failed_cnt} failed")
print(f"Source integrity: {'All sources untouched' if sources_ok else 'SOURCE MODIFIED'}")

# Clean up temporary test files
try:
    shutil.rmtree(TEMP_DIR, ignore_errors=True)
except Exception:
    pass

if failed_cnt > 0 or not sources_ok:
    sys.exit(1)
sys.exit(0)
